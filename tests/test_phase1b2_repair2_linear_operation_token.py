import ast
import io
import tempfile
import unittest
from dataclasses import FrozenInstanceError
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import cv2
import numpy as np
import openpyxl

import yolo_detection_worker as worker_module
from batch_state import (
    RegressionOperationIdentity,
    RegressionSessionState,
    RegressionSet,
)
from detectmain import DetectMain, LINEAR_MODE_SINGLE_IMAGE
from yolo_detection_worker import YoloDetectionWorker


class FakeClock:
    def __init__(self, initial_ns=30_000_000_000):
        self.value_ns = initial_ns
        self.calls = 0

    def __call__(self):
        self.calls += 1
        return self.value_ns

    def advance_ms(self, milliseconds):
        self.value_ns += milliseconds * 1_000_000


class FakeSignal:
    def __init__(self):
        self.values = []

    def emit(self, value):
        self.values.append(value)


class FakeDisplay:
    def __init__(self):
        self.values = []

    def display(self, value):
        self.values.append(value)


class FakeTimer:
    def __init__(self, active=False):
        self.active = active
        self.stop_calls = 0

    def isActive(self):
        return self.active

    def stop(self):
        self.active = False
        self.stop_calls += 1


def formulas(offset=0.0):
    return {
        channel: {
            "slope": index + 1.0 + offset,
            "intercept": index + 0.25,
            "r": 0.9,
            "R2": 0.81,
            "p": 0.01,
            "std_err": 0.02,
        }
        for index, channel in enumerate(("R", "G", "B"))
    }


def regression(source_id="runtime-calibration:C:/calibration/source.png", offset=0.0):
    return RegressionSet.from_formulas(
        formulas(offset),
        source_id=source_id,
        valid_ranges=(1.0, 2.0),
    )


def regression_payload(source_path="C:/calibration/source.png", offset=0.0):
    return {
        "source_path": source_path,
        "image": np.zeros((8, 12, 3), dtype=np.uint8),
        "samples": [
            {
                "No.": 1,
                "Con.": 1.0,
                "Red": 10.0,
                "Green": 20.0,
                "Blue": 30.0,
                "included": True,
            },
            {
                "No.": 2,
                "Con.": 2.0,
                "Red": 11.0,
                "Green": 21.0,
                "Blue": 31.0,
                "included": True,
            },
        ],
        "formulas": formulas(offset),
        "selected_channel": "R",
        "warnings": [],
    }


def accept(session, regression_set, clock, duration_ms):
    operation = session.begin_attempt(regression_set.source_id)
    clock.advance_ms(duration_ms)
    timing = session.accept_success(
        regression_set,
        operation.operation_token,
    )
    if timing is None:
        raise AssertionError("The current operation was unexpectedly rejected.")
    return operation, timing


class RegressionOperationTokenSessionTests(unittest.TestCase):
    def test_operation_identity_is_frozen_and_strictly_validated(self):
        identity = RegressionOperationIdentity("source", 1)
        with self.assertRaises(FrozenInstanceError):
            identity.operation_token = 2
        for invalid in (None, "1", True, False, 0, -1, 1.5):
            with self.subTest(invalid=invalid):
                with self.assertRaises(ValueError):
                    RegressionOperationIdentity("source", invalid)

    def test_rejected_reentry_reads_no_clock_and_consumes_no_token(self):
        clock = FakeClock()
        session = RegressionSessionState(clock_ns=clock)
        first = session.begin_attempt("same-source")
        calls_after_first = clock.calls

        with self.assertRaises(RuntimeError):
            session.begin_attempt("same-source")

        self.assertEqual(clock.calls, calls_after_first)
        self.assertIs(session.active_operation, first)
        session.cancel_attempt()
        second = session.begin_attempt("same-source")
        self.assertEqual(second.operation_token, first.operation_token + 1)

    def test_same_source_stale_success_and_failure_cannot_end_new_operation(self):
        clock = FakeClock()
        session = RegressionSessionState(clock_ns=clock)
        old = regression(offset=0.0)
        operation_one, old_timing = accept(session, old, clock, 11)
        new = regression(offset=10.0)
        operation_two = session.begin_attempt(new.source_id)
        calls_after_new_start = clock.calls

        self.assertIsNone(
            session.accept_success(old, operation_one.operation_token)
        )
        self.assertFalse(
            session.accept_failure(
                operation_one.source_id,
                operation_one.operation_token,
            )
        )
        self.assertEqual(clock.calls, calls_after_new_start)
        self.assertIs(session.active_operation, operation_two)
        self.assertIs(session.active_regression_set, old)
        self.assertIs(session.active_regression_timing, old_timing)

        clock.advance_ms(37)
        timing = session.accept_success(new, operation_two.operation_token)
        self.assertIsNotNone(timing)
        self.assertEqual(timing.duration_ms, 37)
        self.assertEqual(timing.regression_revision, new.revision)
        self.assertIs(session.active_regression_set, new)
        calls_after_success = clock.calls

        self.assertIsNone(
            session.accept_success(new, operation_two.operation_token)
        )
        self.assertFalse(
            session.accept_failure(
                operation_two.source_id,
                operation_two.operation_token,
            )
        )
        self.assertEqual(clock.calls, calls_after_success)
        self.assertIs(session.active_regression_set, new)
        self.assertIs(session.active_regression_timing, timing)

    def test_missing_invalid_wrong_token_and_wrong_source_are_rejected(self):
        clock = FakeClock()
        session = RegressionSessionState(clock_ns=clock)
        current = regression()
        operation = session.begin_attempt(current.source_id)
        calls_after_start = clock.calls

        with self.assertRaises(TypeError):
            session.accept_success(current)
        with self.assertRaises(TypeError):
            session.accept_failure(current.source_id)
        for invalid in (None, "1", True, False, 0, -1, 1.5, operation.operation_token + 1):
            with self.subTest(invalid=invalid):
                self.assertIsNone(session.accept_success(current, invalid))
                self.assertFalse(session.accept_failure(current.source_id, invalid))
        other = regression("runtime-calibration:C:/calibration/other.png")
        self.assertIsNone(
            session.accept_success(other, operation.operation_token)
        )
        self.assertFalse(
            session.accept_failure(other.source_id, operation.operation_token)
        )

        self.assertEqual(clock.calls, calls_after_start)
        self.assertIs(session.active_operation, operation)
        self.assertIsNone(session.active_regression_set)
        self.assertIsNone(session.active_regression_timing)

    def test_tokens_are_never_reused_after_success_failure_cancel_or_import(self):
        clock = FakeClock()
        session = RegressionSessionState(clock_ns=clock)
        current = regression()
        first = session.begin_attempt(current.source_id)
        session.accept_success(current, first.operation_token)

        second = session.begin_attempt(current.source_id)
        self.assertTrue(
            session.accept_failure(second.source_id, second.operation_token)
        )
        session.activate(current)
        third = session.begin_attempt(current.source_id)
        session.cancel_attempt()
        fourth = session.begin_attempt(current.source_id)

        self.assertEqual(
            [
                first.operation_token,
                second.operation_token,
                third.operation_token,
                fourth.operation_token,
            ],
            [1, 2, 3, 4],
        )


class RegressionOperationTokenGuiGateTests(unittest.TestCase):
    def test_same_source_stale_success_and_failure_leave_new_gui_task_busy(self):
        source_path = "C:/calibration/repeated.png"
        source_id = DetectMain._regression_source_id(source_path)
        clock = FakeClock()
        session = RegressionSessionState(clock_ns=clock)
        old = regression(source_id, offset=0.0)
        operation_one, _timing = accept(session, old, clock, 8)
        operation_two = session.begin_attempt(source_id)
        task_changes = []
        harness = SimpleNamespace(
            _active_worker_task="regression",
            _calibration_source_path=source_path,
            _regression_session_state=session,
            _set_active_worker_task=task_changes.append,
        )

        DetectMain._on_regression_finished(
            harness,
            {
                "source_path": source_path,
                "operation_token": operation_one.operation_token,
            },
        )
        DetectMain._on_regression_failed(
            harness,
            {
                "source_path": source_path,
                "operation_token": operation_one.operation_token,
                "message": "late failure",
            },
        )

        self.assertEqual(task_changes, [])
        self.assertEqual(harness._active_worker_task, "regression")
        self.assertIs(session.active_operation, operation_two)
        self.assertTrue(
            DetectMain._linear_regression_result_matches(
                harness,
                {
                    "source_path": source_path,
                    "operation_token": operation_two.operation_token,
                },
            )
        )
        self.assertTrue(
            DetectMain._linear_regression_failure_matches(
                harness,
                {
                    "source_path": source_path,
                    "operation_token": operation_two.operation_token,
                    "message": "current failure",
                },
            )
        )

    def test_gui_gate_rejects_missing_and_wrong_token_types_before_mutation(self):
        source_path = "C:/calibration/current.png"
        session = RegressionSessionState(clock_ns=FakeClock())
        operation = session.begin_attempt(
            DetectMain._regression_source_id(source_path)
        )
        task_changes = []
        harness = SimpleNamespace(
            _active_worker_task="regression",
            _calibration_source_path=source_path,
            _regression_session_state=session,
            _set_active_worker_task=task_changes.append,
        )

        for invalid in (None, "1", True, 0, -1, operation.operation_token + 1):
            with self.subTest(invalid=invalid):
                success = {"source_path": source_path}
                failure = {"source_path": source_path, "message": "failure"}
                if invalid is not None:
                    success["operation_token"] = invalid
                    failure["operation_token"] = invalid
                DetectMain._on_regression_finished(harness, success)
                DetectMain._on_regression_failed(harness, failure)

        self.assertEqual(task_changes, [])
        self.assertEqual(harness._active_worker_task, "regression")
        self.assertIs(session.active_operation, operation)


class RegressionOperationTokenWorkerTests(unittest.TestCase):
    class FakeModel:
        def __init__(self, error=None):
            self.error = error

        def predict(self, **_kwargs):
            if self.error is not None:
                raise self.error
            return [object()]

    class Harness:
        def __init__(self, model):
            self._model_for_test = model
            self._active_formulas = None
            self.regression_finished = FakeSignal()
            self.regression_failed = FakeSignal()

        def _get_model(self, _weight_path):
            return self._model_for_test

        def _build_regression_payload(
            self,
            _result,
            _settings,
            _cv2,
            source_path,
            config_warnings,
        ):
            return {
                "source_path": source_path,
                "formulas": formulas(),
                "warnings": list(config_warnings),
            }

    def invoke(self, model, token):
        harness = self.Harness(model)
        with patch.object(
            worker_module,
            "load_effective_settings",
            return_value=({"detect_confidence": 0.5}, [], None),
        ), patch.object(
            np,
            "fromfile",
            return_value=np.array([1], dtype=np.uint8),
        ), patch.object(
            cv2,
            "imdecode",
            return_value=np.zeros((4, 5, 3), dtype=np.uint8),
        ):
            YoloDetectionWorker.regress(
                harness,
                "C:/calibration/worker.png",
                "C:/weights/best.pt",
                token,
            )
        return harness

    def test_worker_success_and_failure_return_the_original_operation_token(self):
        success = self.invoke(self.FakeModel(), 41)
        self.assertEqual(len(success.regression_finished.values), 1)
        self.assertEqual(success.regression_failed.values, [])
        self.assertEqual(
            success.regression_finished.values[0]["operation_token"],
            41,
        )

        failure = self.invoke(self.FakeModel(RuntimeError("synthetic")), 42)
        self.assertEqual(failure.regression_finished.values, [])
        self.assertEqual(len(failure.regression_failed.values), 1)
        self.assertEqual(
            failure.regression_failed.values[0]["operation_token"],
            42,
        )
        self.assertEqual(
            failure.regression_failed.values[0]["source_path"],
            "C:/calibration/worker.png",
        )

    def test_worker_does_not_replace_an_invalid_received_token(self):
        for invalid in (None, "1", True, 0, -1):
            with self.subTest(invalid=invalid):
                harness = self.invoke(self.FakeModel(), invalid)
                self.assertEqual(harness.regression_finished.values, [])
                self.assertEqual(
                    harness.regression_failed.values[0]["operation_token"],
                    invalid,
                )


class RegressionOperationTokenCompatibilityTests(unittest.TestCase):
    def test_formal_signal_slot_and_callback_signatures_carry_the_token(self):
        project_root = Path(__file__).resolve().parents[1]
        detect_tree = ast.parse(
            (project_root / "detectmain.py").read_text(encoding="utf-8"),
            filename="detectmain.py",
        )
        worker_tree = ast.parse(
            (project_root / "yolo_detection_worker.py").read_text(encoding="utf-8"),
            filename="yolo_detection_worker.py",
        )

        def class_node(tree, name):
            return next(
                node
                for node in tree.body
                if isinstance(node, ast.ClassDef) and node.name == name
            )

        def dotted_name(node):
            if isinstance(node, ast.Name):
                return node.id
            if isinstance(node, ast.Attribute):
                parent = dotted_name(node.value)
                return "{}.{}".format(parent, node.attr) if parent else node.attr
            return None

        detect_class = class_node(detect_tree, "DetectMain")
        worker_class = class_node(worker_tree, "YoloDetectionWorker")
        assignments = {
            target.id: node.value
            for node in detect_class.body + worker_class.body
            if isinstance(node, ast.Assign)
            for target in node.targets
            if isinstance(target, ast.Name)
        }
        self.assertEqual(
            tuple(
                dotted_name(argument)
                for argument in assignments["regression_requested"].args
            ),
            ("str", "str", "object"),
        )
        self.assertEqual(
            tuple(
                dotted_name(argument)
                for argument in assignments["regression_failed"].args
            ),
            ("object",),
        )

        regress = next(
            node
            for node in worker_class.body
            if isinstance(node, ast.FunctionDef) and node.name == "regress"
        )
        self.assertEqual(
            tuple(argument.arg for argument in regress.args.args),
            ("self", "image_path", "weight_path", "operation_token"),
        )
        slot = next(
            decorator
            for decorator in regress.decorator_list
            if isinstance(decorator, ast.Call)
            and dotted_name(decorator.func) == "Slot"
        )
        self.assertEqual(
            tuple(dotted_name(argument) for argument in slot.args),
            ("str", "str", "object"),
        )

        handler_arguments = {
            node.name: tuple(argument.arg for argument in node.args.args)
            for node in detect_class.body
            if isinstance(node, ast.FunctionDef)
            and node.name in {"_on_regression_finished", "_on_regression_failed"}
        }
        self.assertEqual(
            handler_arguments,
            {
                "_on_regression_finished": ("self", "payload"),
                "_on_regression_failed": ("self", "failure"),
            },
        )

    def test_time_display_xlsx_o_and_legacy_loader_contracts_are_unchanged(self):
        payload = regression_payload()
        payload["elapsed_ms"] = 999_000
        current = DetectMain._regression_set_from_result(payload)
        clock = FakeClock()
        session = RegressionSessionState(clock_ns=clock)
        _operation, timing = accept(session, current, clock, 1284)

        display = FakeDisplay()
        display_harness = SimpleNamespace(
            _linear_mode=LINEAR_MODE_SINGLE_IMAGE,
            _calibration_source_path=payload["source_path"],
            _regression_result=payload,
            _regression_session_state=session,
            _linear_duration_timer=FakeTimer(active=True),
            ui=SimpleNamespace(lcdLinearTime=display),
        )
        DetectMain._refresh_linear_duration_display(display_harness)
        self.assertEqual(display.values, ["1284"])
        self.assertEqual(timing.duration_ms, 1284)

        export_harness = SimpleNamespace(
            _regression_result=payload,
            _regression_session_state=session,
            CALIBRATION_TABLE_HEADERS=DetectMain.CALIBRATION_TABLE_HEADERS,
            FORMULA_CHANNELS=DetectMain.FORMULA_CHANNELS,
            FORMULA_FIELDS=DetectMain.FORMULA_FIELDS,
            _validated_source_path=DetectMain._validated_source_path,
            _regression_set_from_result=DetectMain._regression_set_from_result,
        )
        export = DetectMain._validated_linear_export_payload(
            export_harness,
            payload,
        )
        workbook_bytes = DetectMain._build_linear_workbook_bytes(export)
        workbook = openpyxl.load_workbook(io.BytesIO(workbook_bytes))
        try:
            worksheet = workbook["Sheet1"]
            self.assertEqual(worksheet.max_column, 15)
            self.assertEqual(worksheet["O1"].value, "Linear Time (ms)")
            self.assertIsNone(worksheet["G1"].value)
            self.assertEqual(
                tuple(worksheet.cell(1, column).value for column in range(8, 15)),
                ("Channel",) + DetectMain.FORMULA_FIELDS,
            )
            for row in range(2, 4):
                self.assertEqual(worksheet.cell(row, 15).value, 1284)
                self.assertEqual(worksheet.cell(row, 15).number_format, "0")
            self.assertNotIn(
                "operation_token",
                tuple(
                    worksheet.cell(1, column).value
                    for column in range(1, worksheet.max_column + 1)
                ),
            )
            worksheet.delete_cols(15)
            without_time = io.BytesIO()
            workbook.save(without_time)
        finally:
            workbook.close()

        with tempfile.TemporaryDirectory(prefix="phase1b2-repair2-loader-") as directory:
            path = Path(directory) / "linear_con_rgb.xlsx"
            path.write_bytes(without_time.getvalue())
            without_o = DetectMain._load_saved_regression_set(path)
            path.write_bytes(workbook_bytes)
            with_o = DetectMain._load_saved_regression_set(path)

        self.assertEqual(with_o.revision, without_o.revision)
        self.assertEqual(
            with_o.calculate_all(10.0, 20.0, 30.0),
            without_o.calculate_all(10.0, 20.0, 30.0),
        )
        self.assertEqual(
            DetectMain.CALIBRATION_TABLE_HEADERS,
            ("No.", "Con.", "Red", "Green", "Blue"),
        )


if __name__ == "__main__":
    unittest.main()
