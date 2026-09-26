import hashlib
import io
import os
import tempfile
import unittest
from pathlib import Path
from types import MappingProxyType, SimpleNamespace
from unittest.mock import patch

import numpy as np
import openpyxl

import detectmain
from batch_detection_controller import BatchDetectionController
from batch_state import (
    BatchState,
    ConcentrationStatus,
    DetectionScope,
    ImageItem,
    ImageStatus,
    NumberingMode,
    SampleResult,
    assign_batch_numbers,
    detection_table_rows,
    format_detection_duration,
)
from detectmain import DetectMain


class FakeClock:
    def __init__(self, initial_ns=10_000_000_000):
        self.value_ns = initial_ns
        self.calls = 0

    def __call__(self):
        self.calls += 1
        return self.value_ns

    def advance_ms(self, milliseconds):
        self.value_ns += milliseconds * 1_000_000


def lock_context(_path, context):
    return MappingProxyType(dict(context))


def reject_context(_path, _context):
    raise ValueError("synthetic context rejection")


def sample_dict(task, position=1):
    offset = float(position * 20)
    return {
        "image_order": task.image_order,
        "source_file": task.source_file,
        "cuvette_box": (offset, 0.0, offset + 18.0, 40.0),
        "liquid_box": (offset + 3.0, 8.0, offset + 15.0, 35.0),
        "roi_box": (offset + 5.0, 12.0, offset + 13.0, 30.0),
        "red": 10.0 + position,
        "green": 20.0 + position,
        "blue": 30.0 + position,
        "con_r": 0.25 + position,
        "con_g": 0.5 + position,
        "con_b": 0.75 + position,
        "status_r": ConcentrationStatus.IN_RANGE.value,
        "status_g": ConcentrationStatus.BELOW_RANGE.value,
        "status_b": ConcentrationStatus.ABOVE_RANGE.value,
        "no_in_image": position,
        "batch_no": task.batch_start_no + position - 1,
        "status": "valid",
        "warnings": [],
    }


def success_payload(task, sample_count=1):
    return {
        "run_token": task.run_token,
        "job_token": task.job_token,
        "regression_revision": task.regression_revision,
        "image_order": task.image_order,
        "source_path": task.path,
        "source_file": task.source_file,
        "sample_results": [
            sample_dict(task, position)
            for position in range(1, sample_count + 1)
        ],
        "sample_errors": [],
    }


def failure_payload(task, message="synthetic worker failure"):
    return {
        "run_token": task.run_token,
        "job_token": task.job_token,
        "regression_revision": task.regression_revision,
        "image_order": task.image_order,
        "source_path": task.path,
        "source_file": task.source_file,
        "message": message,
    }


def dispatch(controller, task):
    context, rejection = controller.prepare_dispatch_context(task, lock_context)
    if rejection is not None:
        raise AssertionError("A legal test task was rejected.")
    controller.start_dispatch_timer(task)
    return context


def sample(image_order, filename, position):
    offset = float(position * 20)
    return SampleResult(
        image_order=image_order,
        source_file=filename,
        cuvette_box=(offset, 0.0, offset + 18.0, 40.0),
        liquid_box=(offset + 3.0, 8.0, offset + 15.0, 35.0),
        roi_box=(offset + 5.0, 12.0, offset + 13.0, 30.0),
        red=10.0 + position,
        green=20.0 + position,
        blue=30.0 + position,
        con_r=0.25 + position,
        con_g=0.5 + position,
        con_b=0.75 + position,
        status_r=ConcentrationStatus.IN_RANGE,
        status_g=ConcentrationStatus.BELOW_RANGE,
        status_b=ConcentrationStatus.ABOVE_RANGE,
    )


def completed_image(path, image_order, duration_ms, sample_count=2):
    filename = Path(path).name
    image = ImageItem(
        path=str(path),
        original_filename=filename,
        image_order=image_order,
        status=ImageStatus.COMPLETED,
        samples=[sample(image_order, filename, value) for value in range(sample_count)],
        detection_duration_ms=duration_ms,
    )
    return image


class ExportHarness:
    _validated_source_path = staticmethod(DetectMain._validated_source_path)
    _normalized_detection_targets = staticmethod(DetectMain._normalized_detection_targets)
    _normalized_detection_runtime_payload = DetectMain._normalized_detection_runtime_payload

    def __init__(self, image):
        assign_batch_numbers([image])
        state = BatchState(images=[image])
        self._batch_controller = SimpleNamespace(state=state)
        self._detection_cache_state = state
        self._detection_cache_run_token = 73
        self._detection_selected_key = (73, image.image_order, image.path)

    @staticmethod
    def _safe_detection_stem(source_path):
        return DetectMain._safe_detection_stem(source_path)

    def export(self, image, fake_duration=9999):
        targets = []
        for position, authoritative in enumerate(image.samples, start=1):
            targets.append({
                "No.": position,
                "Con.": authoritative.con_r,
                "Red": authoritative.red,
                "Green": authoritative.green,
                "Blue": authoritative.blue,
                "rgb_roi": authoritative.roi_box,
                "detection_duration_ms": fake_duration,
                "Detection Time (ms)": fake_duration,
            })
        payload = {
            "source_path": image.path,
            "image": np.zeros((8, 12, 3), dtype=np.uint8),
            "targets": targets,
            "sample_results": [],
            "sample_errors": [],
            "warnings": [],
            "display_channel": "R",
        }
        return DetectMain._validated_detection_export_payload(self, payload)


class FakeDisplay:
    def __init__(self):
        self.values = []

    def display(self, value):
        self.values.append(value)


class FakeLabel:
    def __init__(self):
        self.values = []

    def setText(self, value):
        self.values.append(value)


class DetectionDurationControllerTests(unittest.TestCase):
    def make_controller(self, paths=("a.png",), scope=DetectionScope.CURRENT_IMAGE):
        clock = FakeClock()
        controller = BatchDetectionController(clock_ns=clock)
        controller.set_options(scope, NumberingMode.PER_IMAGE)
        controller.replace_images(paths)
        return controller, clock

    def test_timer_starts_only_for_prepared_dispatch(self):
        controller, clock = self.make_controller()
        task = controller.begin()

        self.assertIsNone(controller.active_elapsed_ms)
        context, rejection = controller.prepare_dispatch_context(task, lock_context)
        self.assertIsInstance(context, MappingProxyType)
        self.assertIsNone(rejection)
        self.assertIsNone(controller.active_elapsed_ms)
        self.assertEqual(clock.calls, 0)

        controller.start_dispatch_timer(task)
        self.assertEqual(controller.active_elapsed_ms, 0)
        clock.advance_ms(1)
        self.assertEqual(controller.active_elapsed_ms, 1)

    def test_predispatch_context_rejection_has_no_timer_or_duration(self):
        controller, clock = self.make_controller()
        task = controller.begin()
        context, rejection = controller.prepare_dispatch_context(task, reject_context)

        self.assertIsNone(context)
        self.assertIsNotNone(rejection)
        self.assertEqual(clock.calls, 0)
        self.assertIsNone(controller.active_elapsed_ms)
        self.assertIsNone(controller.state.images[0].detection_duration_ms)
        self.assertEqual(controller.state.images[0].status, ImageStatus.FAILED)
        self.assertIsNone(controller.next_task())
        self.assertIsNotNone(controller.finish_if_done())

    def test_success_records_exact_fake_clock_duration(self):
        controller, clock = self.make_controller()
        task = controller.begin()
        dispatch(controller, task)
        clock.advance_ms(1284)

        self.assertTrue(controller.accept_payload(success_payload(task)))
        self.assertEqual(controller.state.images[0].detection_duration_ms, 1284)
        self.assertIsNone(controller.active_elapsed_ms)

    def test_worker_failure_records_duration_without_consuming_batch_number(self):
        controller, clock = self.make_controller(
            ("a.png", "b.png"), DetectionScope.ALL_IMPORTED_IMAGES
        )
        first = controller.begin()
        dispatch(controller, first)
        clock.advance_ms(375)
        self.assertTrue(controller.accept_failure(failure_payload(first)))

        second = controller.next_task()
        self.assertEqual(second.batch_start_no, 1)
        dispatch(controller, second)
        clock.advance_ms(625)
        self.assertTrue(controller.accept_payload(success_payload(second)))
        self.assertEqual(
            [image.detection_duration_ms for image in controller.state.images],
            [375, 625],
        )
        self.assertEqual(controller.state.images[1].samples[0].batch_no, 1)

    def test_wrong_callbacks_leave_active_timer_and_image_unchanged(self):
        controller, clock = self.make_controller()
        task = controller.begin()
        dispatch(controller, task)
        clock.advance_ms(400)
        started_ns = controller._active_started_ns
        original = success_payload(task)
        variants = []
        for field, bad_value in (
            ("run_token", task.run_token + 1),
            ("job_token", task.job_token + 1),
            ("image_order", task.image_order + 1),
            ("regression_revision", "0" * 64),
            ("source_path", task.path + ".wrong"),
            ("source_file", task.source_file + ".wrong"),
        ):
            variant = dict(original)
            variant[field] = bad_value
            variants.append(variant)

        for variant in variants:
            self.assertFalse(controller.accept_payload(variant))
            self.assertFalse(controller.accept_failure(
                dict(variant, message="stale")
            ))
            self.assertEqual(controller._active_started_ns, started_ns)
            self.assertIsNone(controller.state.images[0].detection_duration_ms)

        clock.advance_ms(884)
        self.assertTrue(controller.accept_payload(original))
        self.assertEqual(controller.state.images[0].detection_duration_ms, 1284)

    def test_two_images_receive_independent_durations(self):
        controller, clock = self.make_controller(
            ("a.png", "b.png"), DetectionScope.ALL_IMPORTED_IMAGES
        )
        first = controller.begin()
        dispatch(controller, first)
        clock.advance_ms(120)
        self.assertTrue(controller.accept_payload(success_payload(first)))

        second = controller.next_task()
        self.assertIsNone(controller.active_elapsed_ms)
        dispatch(controller, second)
        clock.advance_ms(987)
        self.assertTrue(controller.accept_payload(success_payload(second)))
        self.assertEqual(
            [image.detection_duration_ms for image in controller.state.images],
            [120, 987],
        )
        self.assertIsNotNone(controller.finish_if_done())
        self.assertIsNone(controller.active_elapsed_ms)

    def test_new_current_image_run_clears_only_planned_duration(self):
        controller, _clock = self.make_controller(("a.png", "b.png"))
        controller.state.images[0].detection_duration_ms = 111
        controller.state.images[1].detection_duration_ms = 222

        controller.begin()

        self.assertIsNone(controller.state.images[0].detection_duration_ms)
        self.assertEqual(controller.state.images[1].detection_duration_ms, 222)
        self.assertIsNone(controller.active_elapsed_ms)

    def test_large_rejection_chains_are_iterative_and_leave_no_timer(self):
        for image_count in (1500, 5000):
            with self.subTest(image_count=image_count):
                controller, clock = self.make_controller(
                    tuple("image-{:04d}.png".format(value) for value in range(image_count)),
                    DetectionScope.ALL_IMPORTED_IMAGES,
                )
                task = controller.begin()
                worker_emits = 0
                while task is not None:
                    context, rejection = controller.prepare_dispatch_context(
                        task, reject_context
                    )
                    self.assertIsNone(context)
                    self.assertIsNotNone(rejection)
                    task = controller.next_task()

                summary = controller.finish_if_done()
                self.assertEqual(worker_emits, 0)
                self.assertEqual(summary["failed_images"], image_count)
                self.assertTrue(all(
                    image.detection_duration_ms is None
                    for image in controller.state.images
                ))
                self.assertIsNone(controller.active_elapsed_ms)
                self.assertEqual(clock.calls, 0)


class DetectionDurationPresentationTests(unittest.TestCase):
    def test_formatter_contract_and_invalid_values(self):
        self.assertEqual(format_detection_duration(None), "")
        self.assertEqual(format_detection_duration(None, empty="—"), "—")
        self.assertEqual(format_detection_duration(0), "0")
        self.assertEqual(format_detection_duration(1), "1")
        self.assertEqual(format_detection_duration(1284), "1284")
        for invalid in (-1, 1.5, True, "1"):
            with self.subTest(invalid=invalid):
                with self.assertRaises(ValueError):
                    format_detection_duration(invalid)

    def test_multi_sample_table_uses_one_image_duration_and_ignores_targets(self):
        image = completed_image("C:/input/shared.png", 1, 1284)
        assign_batch_numbers([image])
        polluted_targets = [
            {"Detection Time (ms)": "9999", "detection_duration_ms": -1}
            for _sample in image.samples
        ]

        rows = detection_table_rows(
            image, NumberingMode.PER_IMAGE, "R"
        )

        self.assertTrue(all(len(row) == 5 for row in rows))
        self.assertEqual([row[-1] for row in rows], [30.0, 31.0])
        self.assertEqual([row[0] for row in rows], [1, 2])
        self.assertEqual(len(polluted_targets), len(rows))
        self.assertEqual(image.detection_duration_ms, 1284)

    def test_browsing_images_displays_each_saved_authoritative_duration(self):
        first = completed_image("C:/input/a.png", 1, 111, sample_count=1)
        second = completed_image("C:/input/b.png", 2, 2345, sample_count=1)
        controller = BatchDetectionController(BatchState(images=[first, second]))
        display = FakeDisplay()
        label = FakeLabel()
        harness = SimpleNamespace(
            _batch_controller=controller,
            _detection_duration_timer=None,
            _detection_selected_key=None,
            ui=SimpleNamespace(label_9=label, lcdNumber=display),
        )

        DetectMain._refresh_detection_duration_display(harness, first)
        DetectMain._refresh_detection_duration_display(harness, second)

        self.assertEqual(display.values, ["111", "2345"])
        self.assertEqual(
            label.values,
            ["Detection Time (ms):", "Detection Time (ms):"],
        )

    def test_ui_source_contract_has_one_initially_generic_time_control(self):
        root = Path(__file__).resolve().parents[1]
        ui_source = (root / "ui" / "detectmain.ui").read_text(encoding="utf-8")
        generated_source = (root / "ui" / "ui_detectmain.py").read_text(encoding="utf-8")

        self.assertIn("Time (ms):", ui_source)
        self.assertIn("Time (ms):", generated_source)
        self.assertNotIn("Detection Time (ms):", ui_source)
        self.assertNotIn("Linear Time (ms):", ui_source)
        self.assertNotIn('name="lcdLinearTime"', ui_source)
        self.assertNotIn("self.lcdLinearTime", generated_source)
        self.assertNotIn("Period:", ui_source)
        self.assertNotIn('setProperty("intValue", 377)', generated_source)

class DetectionDurationWorkbookTests(unittest.TestCase):
    def workbook(self, duration_ms):
        image = completed_image("C:/inputs/=原始 name.png", 1, duration_ms)
        export = ExportHarness(image).export(image, fake_duration=999_000)
        content = DetectMain._build_detection_workbook_bytes(export)
        DetectMain._validate_detection_workbook_bytes(content, export)
        workbook = openpyxl.load_workbook(
            io.BytesIO(content), read_only=False, data_only=False
        )
        return image, export, content, workbook

    def test_u_column_is_numeric_integer_and_a_through_t_are_preserved(self):
        image, export, _content, workbook = self.workbook(1284)
        try:
            worksheet = workbook["Sheet1"]
            self.assertEqual(worksheet.max_column, 21)
            self.assertEqual(
                tuple(worksheet.cell(1, column).value for column in range(1, 22)),
                DetectMain.DETECTION_TABLE_HEADERS,
            )
            self.assertEqual(DetectMain.DETECTION_TABLE_HEADERS[-1], "Detection Time (ms)")
            self.assertIsNone(worksheet["N1"].value)
            for row in range(2, worksheet.max_row + 1):
                self.assertIsNone(worksheet.cell(row, 14).value)
                self.assertEqual(worksheet.cell(row, 21).value, 1284)
                self.assertEqual(worksheet.cell(row, 21).data_type, "n")
                self.assertEqual(worksheet.cell(row, 21).number_format, "0")
                self.assertEqual(worksheet.cell(row, 2).value, image.original_filename)
                self.assertEqual(worksheet.cell(row, 2).data_type, "s")
            self.assertEqual(export["detection_duration_ms"], 1284)
            self.assertNotEqual(export["detection_duration_ms"], 999_000)
        finally:
            workbook.close()

    def test_none_duration_leaves_u_cells_empty(self):
        _image, export, _content, workbook = self.workbook(None)
        try:
            worksheet = workbook["Sheet1"]
            for row in range(2, worksheet.max_row + 1):
                self.assertIsNone(worksheet.cell(row, 21).value)
                self.assertEqual(worksheet.cell(row, 21).number_format, "0")
            self.assertIsNone(export["detection_duration_ms"])
        finally:
            workbook.close()

    def test_pair_transaction_suffix_rollback_and_input_file_protection(self):
        with tempfile.TemporaryDirectory(prefix="phase1b2-duration-") as temporary_directory:
            root = Path(temporary_directory)
            source = root / "原始 input.png"
            source.write_bytes(b"immutable-original-input")
            before = source.stat()
            before_hash = hashlib.sha256(source.read_bytes()).hexdigest()
            target = root / "saved"

            first = DetectMain._commit_detection_export(
                target, "result", b"xlsx-bytes", b"png-bytes"
            )
            second = DetectMain._commit_detection_export(
                target, "result", b"xlsx-bytes", b"png-bytes"
            )
            self.assertTrue(first.committed)
            self.assertTrue(second.committed)
            self.assertEqual(
                [[path.name for path in result.final_paths] for result in (first, second)],
                [["result.xlsx", "result.png"], ["result_1.xlsx", "result_1.png"]],
            )

            call_count = 0
            real_link = os.link

            def fail_second_publish(source_path, destination):
                nonlocal call_count
                call_count += 1
                if call_count == 2:
                    raise OSError("synthetic PNG publish failure")
                return real_link(source_path, destination)

            with patch.object(detectmain.os, "link", side_effect=fail_second_publish):
                rollback = DetectMain._commit_detection_export(
                    target, "rollback", b"xlsx-bytes", b"png-bytes"
                )
            self.assertFalse(rollback.committed)
            self.assertTrue(rollback.rollback_complete)
            self.assertFalse((target / "rollback.xlsx").exists())
            self.assertFalse((target / "rollback.png").exists())

            after = source.stat()
            self.assertEqual(source.read_bytes(), b"immutable-original-input")
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before_hash)
            self.assertEqual(after.st_size, before.st_size)
            self.assertEqual(after.st_mtime_ns, before.st_mtime_ns)


if __name__ == "__main__":
    unittest.main()
