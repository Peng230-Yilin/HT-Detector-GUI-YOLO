import hashlib
import io
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import numpy as np
import openpyxl

import detectmain
from batch_state import (
    BatchState,
    ConcentrationStatus,
    ImageItem,
    ImageStatus,
    NumberingMode,
    SampleResult,
    assign_batch_numbers,
)
from detectmain import DetectMain


def make_sample(image_order, filename, position):
    x0 = float(position * 20)
    return SampleResult(
        image_order=image_order,
        source_file=filename,
        cuvette_box=(x0, 0.0, x0 + 18.0, 40.0),
        liquid_box=(x0 + 3.0, 8.0, x0 + 15.0, 35.0),
        roi_box=(x0 + 5.0, 12.0, x0 + 13.0, 30.0),
        red=10.0 + position,
        green=20.0 + position,
        blue=30.0 + position,
        con_r=0.25 + position,
        con_g=0.5 + position,
        con_b=0.75 + position,
        status_r=ConcentrationStatus.IN_RANGE,
        status_g=ConcentrationStatus.BELOW_RANGE,
        status_b=ConcentrationStatus.ABOVE_RANGE,
    )


def make_image(path, image_order, sample_count, status=ImageStatus.COMPLETED):
    filename = Path(path).name
    samples = [
        make_sample(image_order, filename, position)
        for position in range(sample_count)
    ]
    return ImageItem(
        path=str(path),
        original_filename=filename,
        image_order=image_order,
        status=status,
        samples=samples,
    )


def runtime_payload(image, display_numbers=None):
    numbers = display_numbers or list(range(1, len(image.samples) + 1))
    targets = []
    for index, (sample, display_number) in enumerate(zip(image.samples, numbers), start=1):
        targets.append({
            "No.": display_number,
            "Con.": index / 4.0,
            "Red": sample.red,
            "Green": sample.green,
            "Blue": sample.blue,
            "rgb_roi": sample.roi_box,
        })
    return {
        "source_path": image.path,
        "image": np.full((8, 12, 3), image.image_order, dtype=np.uint8),
        "targets": targets,
        "display_channel": "R",
    }


class ExportHarness:
    _validated_source_path = staticmethod(DetectMain._validated_source_path)
    _normalized_detection_targets = staticmethod(DetectMain._normalized_detection_targets)
    _normalized_detection_runtime_payload = DetectMain._normalized_detection_runtime_payload

    def __init__(self, state):
        self._batch_controller = SimpleNamespace(state=state)
        self._detection_cache_state = state
        self._detection_cache_run_token = 41
        self._detection_selected_key = None

    @staticmethod
    def _safe_detection_stem(source_path):
        return DetectMain._safe_detection_stem(source_path)

    def export(self, image, display_numbers=None):
        self._detection_selected_key = (
            self._detection_cache_run_token,
            image.image_order,
            image.path,
        )
        return DetectMain._validated_detection_export_payload(
            self, runtime_payload(image, display_numbers)
        )


def workbook_rows(export_payload):
    content = DetectMain._build_detection_workbook_bytes(export_payload)
    DetectMain._validate_detection_workbook_bytes(content, export_payload)
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        worksheet = workbook["Sheet1"]
        rows = tuple(
            tuple(worksheet.cell(row, column).value for column in range(1, 22))
            for row in range(1, worksheet.max_row + 1)
        )
    finally:
        workbook.close()
    return content, rows


class DetectionXlsxSchemeBTests(unittest.TestCase):
    def state_with_numbers(self, images, mode=NumberingMode.PER_IMAGE):
        assign_batch_numbers(images)
        return BatchState(images=images, numbering_mode=mode)

    def test_single_image_single_sample_has_scheme_b_then_all_v1_columns(self):
        image = make_image("C:/inputs/sample one.png", 1, 1)
        state = self.state_with_numbers([image])
        export = ExportHarness(state).export(image, display_numbers=[99])
        _, rows = workbook_rows(export)

        self.assertEqual(rows[0], DetectMain.DETECTION_TABLE_HEADERS)
        self.assertEqual(rows[1], (
            1, "sample one.png", 1, 1,
            10, 20, 30, 0.25, 0.5, 0.75,
            "IN_RANGE", "BELOW_RANGE", "ABOVE_RANGE", None,
            5, 12, 13, 30, 8, 18,
            None,
        ))

    def test_single_image_multiple_samples_restart_no_in_image_at_one(self):
        image = make_image("C:/inputs/many samples.jpeg", 1, 3)
        state = self.state_with_numbers([image])
        export = ExportHarness(state).export(image, display_numbers=[8, 9, 10])
        _, rows = workbook_rows(export)

        self.assertEqual([row[0] for row in rows[1:]], [1, 1, 1])
        self.assertEqual([row[2] for row in rows[1:]], [1, 2, 3])
        self.assertEqual([row[3] for row in rows[1:]], [1, 2, 3])

    def test_selected_image_only_keeps_batch_identity_across_multiple_images(self):
        first = make_image("C:/batch/first.png", 1, 2)
        second = make_image("C:/batch/second.png", 2, 2)
        state = self.state_with_numbers([first, second])
        export = ExportHarness(state).export(second)
        _, rows = workbook_rows(export)

        self.assertEqual(len(rows), 3)
        self.assertEqual([row[:4] for row in rows[1:]], [
            (2, "second.png", 1, 3),
            (2, "second.png", 2, 4),
        ])

    def test_failed_image_does_not_consume_batch_number(self):
        first = make_image("C:/batch/first.png", 1, 2)
        failed = make_image("C:/batch/failed.png", 2, 0, ImageStatus.FAILED)
        third = make_image("C:/batch/third.png", 3, 2)
        state = self.state_with_numbers([first, failed, third])
        export = ExportHarness(state).export(third)
        _, rows = workbook_rows(export)

        self.assertEqual([row[:4] for row in rows[1:]], [
            (3, "third.png", 1, 3),
            (3, "third.png", 2, 4),
        ])

    def test_original_filename_preserves_unicode_spaces_punctuation_and_formula_prefix(self):
        filename = "=中文 English 样品，批次 #1! (最终).PNG"
        image = make_image("C:/输入目录/" + filename, 1, 1)
        state = self.state_with_numbers([image])
        export = ExportHarness(state).export(image)
        _, rows = workbook_rows(export)

        self.assertEqual(rows[1][1], filename)
        workbook_bytes = DetectMain._build_detection_workbook_bytes(export)
        workbook = openpyxl.load_workbook(io.BytesIO(workbook_bytes), data_only=False)
        try:
            source_cell = workbook["Sheet1"]["B2"]
            self.assertEqual(source_cell.value, filename)
            self.assertEqual(source_cell.data_type, "s")
        finally:
            workbook.close()

    def test_same_named_images_are_distinguished_by_image_order(self):
        first = make_image("C:/batch/a/same name.png", 1, 1)
        second = make_image("C:/batch/b/same name.png", 2, 1)
        state = self.state_with_numbers([first, second])
        harness = ExportHarness(state)
        first_row = workbook_rows(harness.export(first))[1][1]
        second_row = workbook_rows(harness.export(second))[1][1]

        self.assertEqual(first_row[:4], (1, "same name.png", 1, 1))
        self.assertEqual(second_row[:4], (2, "same name.png", 1, 2))

    def test_legacy_display_numbers_and_fake_concentration_cannot_change_export(self):
        image = make_image("C:/batch/mode.png", 2, 2)
        for sample, no_in_image, batch_no in zip(image.samples, (1, 2), (5, 6)):
            sample.no_in_image = no_in_image
            sample.batch_no = batch_no
        state = BatchState(images=[image], numbering_mode=NumberingMode.PER_IMAGE)
        harness = ExportHarness(state)

        per_image = harness.export(image, display_numbers=[1, 2])
        state.numbering_mode = NumberingMode.CONTINUOUS_BATCH
        continuous = harness.export(image, display_numbers=[5, 6])

        self.assertEqual(
            [tuple(target[key] for key in (
                "image_order", "source_file", "no_in_image", "batch_no"
            )) for target in per_image["targets"]],
            [tuple(target[key] for key in (
                "image_order", "source_file", "no_in_image", "batch_no"
            )) for target in continuous["targets"]],
        )
        self.assertEqual(
            [target["Con.R"] for target in per_image["targets"]],
            [sample.con_r for sample in image.samples],
        )
        self.assertEqual(
            [target["Con.R"] for target in continuous["targets"]],
            [sample.con_r for sample in image.samples],
        )

    def test_all_formula_prefixes_are_exact_non_formula_strings(self):
        for prefix in ("=", "+", "-", "@"):
            with self.subTest(prefix=prefix):
                filename = prefix + "sample name.PNG"
                image = make_image("C:/inputs/" + filename, 1, 1)
                state = self.state_with_numbers([image])
                export = ExportHarness(state).export(image)
                content = DetectMain._build_detection_workbook_bytes(export)
                DetectMain._validate_detection_workbook_bytes(content, export)
                workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
                try:
                    cell = workbook["Sheet1"]["B2"]
                    self.assertEqual(cell.value, filename)
                    self.assertEqual(cell.data_type, "s")
                finally:
                    workbook.close()

    def test_invalid_channel_concentration_is_blank_and_status_is_preserved(self):
        image = make_image("C:/inputs/partial.png", 1, 1)
        image.samples[0].con_g = None
        image.samples[0].status_g = ConcentrationStatus.INVALID_SLOPE
        state = self.state_with_numbers([image])
        export = ExportHarness(state).export(image)
        _, rows = workbook_rows(export)
        self.assertIsNone(rows[1][8])
        self.assertEqual(rows[1][11], "INVALID_SLOPE")

    def test_pair_transaction_suffixes_without_overwriting_prior_results(self):
        image = make_image("C:/batch/result.png", 1, 1)
        state = self.state_with_numbers([image])
        export = ExportHarness(state).export(image)
        workbook_bytes, png_bytes = DetectMain._build_detection_export_bytes(export)

        with tempfile.TemporaryDirectory() as temporary_directory:
            target = Path(temporary_directory) / "output"
            results = [
                DetectMain._commit_detection_export(
                    target, export["safe_stem"], workbook_bytes, png_bytes
                )
                for _ in range(3)
            ]
            self.assertTrue(all(result.committed for result in results))
            self.assertEqual(
                [[path.name for path in result.final_paths] for result in results],
                [
                    ["result.xlsx", "result.png"],
                    ["result_1.xlsx", "result_1.png"],
                    ["result_2.xlsx", "result_2.png"],
                ],
            )
            for result in results:
                self.assertEqual(result.final_paths[0].read_bytes(), workbook_bytes)
                self.assertEqual(result.final_paths[1].read_bytes(), png_bytes)
            workbook = openpyxl.load_workbook(
                results[0].final_paths[0], read_only=False, data_only=False
            )
            try:
                worksheet = workbook["Sheet1"]
                self.assertEqual(
                    tuple(worksheet.cell(1, column).value for column in range(1, 22)),
                    DetectMain.DETECTION_TABLE_HEADERS,
                )
                self.assertTrue(
                    all(
                        worksheet.cell(row, 14).value is None
                        for row in range(1, worksheet.max_row + 1)
                    )
                )
            finally:
                workbook.close()

    def test_pair_transaction_rolls_back_when_second_publish_fails(self):
        image = make_image("C:/batch/rollback.png", 1, 1)
        state = self.state_with_numbers([image])
        export = ExportHarness(state).export(image)
        workbook_bytes, png_bytes = DetectMain._build_detection_export_bytes(export)
        real_link = os.link
        call_count = 0

        def fail_png_publish(source, destination):
            nonlocal call_count
            call_count += 1
            if call_count == 2:
                raise OSError("synthetic PNG publish failure")
            return real_link(source, destination)

        with tempfile.TemporaryDirectory() as temporary_directory:
            target = Path(temporary_directory) / "output"
            with patch.object(detectmain.os, "link", side_effect=fail_png_publish):
                result = DetectMain._commit_detection_export(
                    target, export["safe_stem"], workbook_bytes, png_bytes
                )
            self.assertFalse(result.committed)
            self.assertTrue(result.rollback_complete)
            self.assertEqual(list(target.iterdir()), [])

    def test_export_and_save_do_not_modify_original_input_file(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            source = root / "原始 input !.png"
            original_bytes = b"immutable-original-input"
            source.write_bytes(original_bytes)
            original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
            original_stat = source.stat()

            image = make_image(source, 1, 1)
            state = self.state_with_numbers([image])
            export = ExportHarness(state).export(image)
            workbook_bytes, png_bytes = DetectMain._build_detection_export_bytes(export)
            result = DetectMain._commit_detection_export(
                root / "saved", export["safe_stem"], workbook_bytes, png_bytes
            )

            self.assertTrue(result.committed)
            self.assertEqual(source.read_bytes(), original_bytes)
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original_hash)
            self.assertEqual(source.stat().st_size, original_stat.st_size)
            self.assertEqual(source.stat().st_mtime_ns, original_stat.st_mtime_ns)


if __name__ == "__main__":
    unittest.main()
