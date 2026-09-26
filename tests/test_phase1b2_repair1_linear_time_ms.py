import io
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import openpyxl

import detectmain
from batch_detection_controller import BatchDetectionController
from batch_state import (
    BatchState,
    ImageItem,
    ImageStatus,
    RegressionSessionState,
    RegressionSet,
    RegressionTiming,
)
from detectmain import DetectMain, LINEAR_MODE_SINGLE_IMAGE


class FakeClock:
    def __init__(self, initial_ns=20_000_000_000):
        self.value_ns = initial_ns
        self.calls = 0

    def __call__(self):
        self.calls += 1
        return self.value_ns

    def advance_ms(self, milliseconds):
        self.value_ns += milliseconds * 1_000_000


class FakeDisplay:
    def __init__(self):
        self.values = []

    def display(self, value):
        self.values.append(value)


class FakeLabel:
    def __init__(self):
        self.values = []

    def setText(self, value):
        self.values.append(value)


class FakeTimer:
    def __init__(self, active=False):
        self.active = active
        self.stop_calls = 0

    def isActive(self):
        return self.active

    def stop(self):
        self.active = False
        self.stop_calls += 1


def formulas(offset=0.0):
    return {
        channel: {
            "slope": index + 1.0 + offset,
            "intercept": index + 0.25,
            "r": 0.9,
            "R2": 0.81,
            "p": 0.01,
            "std_err": 0.02,
        }
        for index, channel in enumerate(("R", "G", "B"))
    }


def regression(source_path="C:/calibration/source.png", offset=0.0):
    return RegressionSet.from_formulas(
        formulas(offset),
        source_id=DetectMain._regression_source_id(source_path),
        valid_ranges=(1.0, 2.0),
    )


def regression_payload(source_path="C:/calibration/source.png", offset=0.0):
    return {
        "source_path": source_path,
        "image": np.zeros((8, 12, 3), dtype=np.uint8),
        "samples": [
            {
                "No.": 1,
                "Con.": 1.0,
                "Red": 10.0,
                "Green": 20.0,
                "Blue": 30.0,
                "included": True,
            },
            {
                "No.": 2,
                "Con.": 2.0,
                "Red": 11.0,
                "Green": 21.0,
                "Blue": 31.0,
                "included": True,
            },
        ],
        "formulas": formulas(offset),
        "selected_channel": "R",
        "warnings": [],
    }


def export_harness(payload, session):
    return SimpleNamespace(
        _regression_result=payload,
        _regression_session_state=session,
        CALIBRATION_TABLE_HEADERS=DetectMain.CALIBRATION_TABLE_HEADERS,
        FORMULA_CHANNELS=DetectMain.FORMULA_CHANNELS,
        FORMULA_FIELDS=DetectMain.FORMULA_FIELDS,
        _validated_source_path=DetectMain._validated_source_path,
        _regression_set_from_result=DetectMain._regression_set_from_result,
    )


def accept_duration(session, regression_set, clock, duration_ms):
    operation = session.begin_attempt(regression_set.source_id)
    clock.advance_ms(duration_ms)
    timing = session.accept_success(
        regression_set,
        operation.operation_token,
    )
    if timing is None:
        raise AssertionError("A valid test regression was rejected.")
    return timing


class LinearTimingAuthorityTests(unittest.TestCase):
    def test_invalid_start_does_not_read_clock(self):
        clock = FakeClock()
        session = RegressionSessionState(clock_ns=clock)

        for invalid in (None, "", 7, True):
            with self.subTest(invalid=invalid):
                with self.assertRaises(ValueError):
                    session.begin_attempt(invalid)
        self.assertEqual(clock.calls, 0)
        self.assertFalse(session.attempt_active)

    def test_success_records_exact_integer_milliseconds_and_identity(self):
        clock = FakeClock()
        session = RegressionSessionState(clock_ns=clock)
        current = regression()

        timing = accept_duration(session, current, clock, 1284)

        self.assertEqual(timing.duration_ms, 1284)
        self.assertEqual(timing.source_id, current.source_id)
        self.assertEqual(timing.regression_revision, current.revision)
        self.assertIs(session.active_regression_set, current)
        self.assertIs(session.timing_for(current), timing)
        self.assertFalse(session.attempt_active)
        self.assertEqual(clock.calls, 2)

        zero_clock = FakeClock()
        zero_session = RegressionSessionState(clock_ns=zero_clock)
        zero_timing = accept_duration(
            zero_session, regression(), zero_clock, 0
        )
        self.assertEqual(zero_timing.duration_ms, 0)

    def test_repeated_regression_replaces_time_without_changing_revision(self):
        clock = FakeClock()
        session = RegressionSessionState(clock_ns=clock)
        first = regression()
        first_timing = accept_duration(session, first, clock, 120)
        second = regression()

        second_timing = accept_duration(session, second, clock, 987)

        self.assertEqual(first.revision, second.revision)
        self.assertEqual(first_timing.duration_ms, 120)
        self.assertEqual(second_timing.duration_ms, 987)
        self.assertIs(session.timing_for(second), second_timing)

    def test_stale_source_and_duplicate_success_do_not_read_end_clock(self):
        clock = FakeClock()
        session = RegressionSessionState(clock_ns=clock)
        current = regression("C:/calibration/current.png")
        stale = regression("C:/calibration/stale.png")
        operation = session.begin_attempt(current.source_id)
        calls_after_start = clock.calls

        self.assertIsNone(
            session.accept_success(stale, operation.operation_token)
        )
        self.assertEqual(clock.calls, calls_after_start)
        self.assertTrue(session.attempt_active)

        clock.advance_ms(500)
        self.assertIsNotNone(
            session.accept_success(current, operation.operation_token)
        )
        calls_after_success = clock.calls
        self.assertIsNone(
            session.accept_success(current, operation.operation_token)
        )
        self.assertEqual(clock.calls, calls_after_success)

    def test_failure_preserves_old_revision_and_timing(self):
        clock = FakeClock()
        session = RegressionSessionState(clock_ns=clock)
        old = regression(offset=0.0)
        old_timing = accept_duration(session, old, clock, 321)
        operation = session.begin_attempt(old.source_id)
        clock.advance_ms(400)
        calls_before_failure = clock.calls

        self.assertIs(
            session.record_failure(
                operation.source_id,
                operation.operation_token,
            ),
            old,
        )

        self.assertEqual(clock.calls, calls_before_failure)
        self.assertIs(session.active_regression_set, old)
        self.assertIs(session.timing_for(old), old_timing)
        self.assertFalse(session.attempt_active)

    def test_failure_without_old_success_has_no_authoritative_time(self):
        clock = FakeClock()
        session = RegressionSessionState(clock_ns=clock)
        source_id = regression().source_id
        operation = session.begin_attempt(source_id)

        self.assertIsNone(
            session.record_failure(
                operation.source_id,
                operation.operation_token,
            )
        )
        self.assertIsNone(session.active_regression_timing)
        self.assertFalse(session.attempt_active)

    def test_timing_validation_rejects_invalid_duration(self):
        current = regression()
        for invalid in (-1, 1.5, True, "1"):
            with self.subTest(invalid=invalid):
                with self.assertRaises(ValueError):
                    RegressionTiming(
                        current.source_id, current.revision, invalid
                    )
        self.assertIsNone(
            RegressionTiming(current.source_id, current.revision, None).duration_ms
        )


class LinearTimingPresentationTests(unittest.TestCase):
    def make_display_harness(self, payload, session, source_path=None):
        return SimpleNamespace(
            _linear_mode=LINEAR_MODE_SINGLE_IMAGE,
            _calibration_source_path=(
                payload["source_path"] if source_path is None else source_path
            ),
            _regression_result=payload,
            _regression_session_state=session,
            _linear_duration_timer=FakeTimer(active=True),
            ui=SimpleNamespace(label_9=FakeLabel(), lcdNumber=FakeDisplay()),
        )

    def test_display_uses_active_then_fixed_authoritative_milliseconds(self):
        payload = regression_payload()
        current = DetectMain._regression_set_from_result(payload)
        clock = FakeClock()
        session = RegressionSessionState(clock_ns=clock)
        harness = self.make_display_harness(payload, session)
        operation = session.begin_attempt(current.source_id)
        clock.advance_ms(77)

        DetectMain._refresh_linear_duration_display(harness)
        session.accept_success(current, operation.operation_token)
        DetectMain._refresh_linear_duration_display(harness)

        self.assertEqual(harness.ui.lcdNumber.values, ["77", "77"])
        self.assertEqual(
            harness.ui.label_9.values,
            ["Linear Time (ms):", "Linear Time (ms):"],
        )
        self.assertEqual(harness._linear_duration_timer.stop_calls, 1)

    def test_new_or_mismatched_source_displays_blank(self):
        payload = regression_payload()
        current = DetectMain._regression_set_from_result(payload)
        clock = FakeClock()
        session = RegressionSessionState(clock_ns=clock)
        accept_duration(session, current, clock, 222)
        harness = self.make_display_harness(
            payload, session, source_path="C:/calibration/new.png"
        )

        DetectMain._refresh_linear_duration_display(harness)

        self.assertEqual(harness.ui.lcdNumber.values, [""])
        self.assertEqual(harness.ui.label_9.values, ["Linear Time (ms):"])

    def test_detection_and_linear_displays_read_independent_authorities(self):
        payload = regression_payload()
        current = DetectMain._regression_set_from_result(payload)
        linear_clock = FakeClock()
        session = RegressionSessionState(clock_ns=linear_clock)
        accept_duration(session, current, linear_clock, 456)
        image = ImageItem(
            path="C:/detection/image.png",
            original_filename="image.png",
            image_order=1,
            status=ImageStatus.COMPLETED,
            detection_duration_ms=789,
        )
        controller = BatchDetectionController(BatchState(images=[image]))
        shared_display = FakeDisplay()
        shared_label = FakeLabel()
        harness = SimpleNamespace(
            _batch_controller=controller,
            _detection_duration_timer=None,
            _detection_selected_key=None,
            _linear_mode=LINEAR_MODE_SINGLE_IMAGE,
            _calibration_source_path=payload["source_path"],
            _regression_result=payload,
            _regression_session_state=session,
            _linear_duration_timer=FakeTimer(),
            ui=SimpleNamespace(
                label_9=shared_label,
                lcdNumber=shared_display,
            ),
        )

        DetectMain._refresh_detection_duration_display(harness, image)
        DetectMain._refresh_linear_duration_display(harness)

        self.assertEqual(shared_display.values, ["789", "456"])
        self.assertEqual(
            shared_label.values,
            ["Detection Time (ms):", "Linear Time (ms):"],
        )

    def test_ui_sources_have_one_shared_control_and_no_gui_time_columns(self):
        root = Path(__file__).resolve().parents[1]
        ui_source = (root / "ui" / "detectmain.ui").read_text(encoding="utf-8")
        generated = (root / "ui" / "ui_detectmain.py").read_text(encoding="utf-8")
        main_source = (root / "detectmain.py").read_text(encoding="utf-8")

        self.assertIn('name="lcdNumber"', ui_source)
        self.assertNotIn('name="lcdLinearTime"', ui_source)
        self.assertNotIn('name="labelLinearTime"', ui_source)
        self.assertNotIn("self.lcdLinearTime", generated)
        self.assertNotIn("self.labelLinearTime", generated)
        self.assertIn("Time (ms):", ui_source)
        self.assertEqual(
            DetectMain.CALIBRATION_TABLE_HEADERS,
            ("No.", "Con.", "Red", "Green", "Blue"),
        )
        self.assertNotIn('"Time (ms)"', main_source)


class LinearTimingWorkbookTests(unittest.TestCase):
    def make_export(self, duration_ms=1284, matching=True):
        payload = regression_payload()
        payload["elapsed_ms"] = 999_000
        current = DetectMain._regression_set_from_result(payload)
        clock = FakeClock()
        session = RegressionSessionState(clock_ns=clock)
        if matching:
            accept_duration(session, current, clock, duration_ms)
        else:
            other = regression("C:/calibration/other.png")
            accept_duration(session, other, clock, duration_ms)
        harness = export_harness(payload, session)
        export = DetectMain._validated_linear_export_payload(harness, payload)
        content = DetectMain._build_linear_workbook_bytes(export)
        DetectMain._validate_linear_workbook_bytes(content, export)
        workbook = openpyxl.load_workbook(
            io.BytesIO(content), read_only=False, data_only=False
        )
        return current, export, workbook

    def test_linear_xlsx_appends_numeric_integer_time_without_moving_formulas(self):
        current, export, workbook = self.make_export(1284)
        try:
            worksheet = workbook["Sheet1"]
            self.assertEqual(worksheet.max_column, 15)
            self.assertEqual(worksheet["O1"].value, "Linear Time (ms)")
            self.assertEqual(
                tuple(worksheet.cell(1, column).value for column in range(8, 15)),
                ("Channel",) + DetectMain.FORMULA_FIELDS,
            )
            self.assertIsNone(worksheet["G1"].value)
            for row in range(2, 4):
                self.assertEqual(worksheet.cell(row, 15).value, 1284)
                self.assertEqual(worksheet.cell(row, 15).data_type, "n")
                self.assertEqual(worksheet.cell(row, 15).number_format, "0")
            self.assertEqual(export["linear_duration_ms"], 1284)
            self.assertNotEqual(export["linear_duration_ms"], 999_000)
            self.assertEqual(export["regression_set"].revision, current.revision)
        finally:
            workbook.close()

    def test_source_mismatch_keeps_linear_xlsx_time_cells_blank(self):
        _current, export, workbook = self.make_export(500, matching=False)
        try:
            worksheet = workbook["Sheet1"]
            for row in range(2, 4):
                self.assertIsNone(worksheet.cell(row, 15).value)
                self.assertEqual(worksheet.cell(row, 15).number_format, "0")
            self.assertIsNone(export["linear_duration_ms"])
        finally:
            workbook.close()

    def test_time_does_not_change_revision_or_three_channel_calculation(self):
        current, export, workbook = self.make_export(999)
        try:
            rebuilt = RegressionSet.from_formulas(
                formulas(), current.source_id, valid_ranges=(1.0, 2.0)
            )
            self.assertEqual(rebuilt.revision, current.revision)
            self.assertEqual(
                rebuilt.calculate_all(10.0, 20.0, 30.0),
                current.calculate_all(10.0, 20.0, 30.0),
            )
            self.assertEqual(
                tuple(workbook["Sheet1"].cell(1, column).value for column in range(8, 15)),
                ("Channel",) + DetectMain.FORMULA_FIELDS,
            )
        finally:
            workbook.close()

        workbook_bytes = DetectMain._build_linear_workbook_bytes(export)
        with tempfile.TemporaryDirectory(prefix="phase1b2-linear-loader-") as directory:
            formula_path = Path(directory) / "linear_con_rgb.xlsx"
            formula_path.write_bytes(workbook_bytes)
            loaded = DetectMain._load_saved_regression_set(formula_path)
            expected = RegressionSet.from_formulas(
                formulas(),
                source_id="saved-linear:{}".format(str(formula_path)),
                valid_ranges=(1.0, 2.0),
            )
            self.assertEqual(loaded.revision, expected.revision)
            self.assertEqual(
                loaded.calculate_all(10.0, 20.0, 30.0),
                expected.calculate_all(10.0, 20.0, 30.0),
            )


if __name__ == "__main__":
    unittest.main()
